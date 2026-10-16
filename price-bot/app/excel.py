from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .db import DB

HEADER = ["Категория", "Раздел", "Товар", "Примечание", "Закупка, ₽", "Наценка, ₽", "Цена AVIOR, ₽"]
WIDTHS = [22, 30, 48, 30, 13, 12, 15]


def build_xlsx(db: DB, target) -> int:
    """Пишет текущий прайс бота в Excel (путь или BytesIO). Возвращает число позиций."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Прайс AVIOR"
    ws.append(HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
        cell.alignment = Alignment(vertical="center")
    for col, width in zip("ABCDEFG", WIDTHS):
        ws.column_dimensions[col].width = width

    rows = 0
    for cat in db.categories():
        for it in db.items(cat["id"]):
            sell = db.sell_price(it, cat)
            ws.append([cat["title"], it["section"], it["name"], it["note"], it["base_price"], sell - it["base_price"], sell])
            rows += 1
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=7):
        for cell in row:
            cell.number_format = "#,##0"
    for cell in ws["G"][1:]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(target)
    return rows
